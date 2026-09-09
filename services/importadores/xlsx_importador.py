import io

from openpyxl import load_workbook


class XlsxImportador:

  @staticmethod
  def importar(conteudo: bytes):

    workbook = load_workbook(io.BytesIO(conteudo), data_only=True)
    planilha = workbook.active
    linhas = list(planilha.iter_rows(values_only=True))
    cabecalho = linhas[0]
    registros = []

    for linha in linhas[1:]:

      registro = {}

      for coluna, valor in zip(cabecalho,linha):

        if (valor is None):
          valor = ""
        else:
          valor = str(valor)

        registro[coluna] = valor

      registros.append(registro)

    return registros

