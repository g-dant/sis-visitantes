from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ResultadoImportacaoXlsxService:

    @staticmethod
    def gerar(linhas: list[dict[str, Any]]) -> BytesIO:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Importação"

        dados_linhas = [
            ResultadoImportacaoXlsxService.__obter_dados_linha(linha)
            for linha in linhas]

        headers = ResultadoImportacaoXlsxService.__obter_headers(dados_linhas)

        lado = Side(
            style="thin",
            color="000000")

        borda = Border(
            left=lado,
            right=lado,
            top=lado,
            bottom=lado)

        alinhamento = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True)

        preenchimento_header = PatternFill(
            fill_type="solid",
            fgColor="000000")

        preenchimento_erro = PatternFill(
            fill_type="solid",
            fgColor="FCE4E4")

        fonte_header = Font(
            color="FFFFFF",
            bold=True)

        # Cabeçalho
        for numero_coluna, header in enumerate(headers, start=1):
            celula = worksheet.cell(
                row=1,
                column=numero_coluna,
                value=header)

            celula.fill = preenchimento_header
            celula.font = fonte_header
            celula.border = borda
            celula.alignment = alinhamento

        # Linhas
        for numero_linha, linha in enumerate(linhas, start=2):
            dados = ResultadoImportacaoXlsxService.__obter_dados_linha(
                linha)

            resultado = str(dados.get("Resultado") or "")
            possui_erro = bool(resultado.strip())

            for numero_coluna, header in enumerate(headers, start=1):

              for numero_coluna, header in enumerate(headers, start=1):
                valor = dados.get(header, "")

                celula = worksheet.cell(
                    row=numero_linha,
                    column=numero_coluna,
                    value=valor)

                celula.border = borda
                celula.alignment = alinhamento

                if possui_erro:
                    celula.fill = preenchimento_erro

        if worksheet.max_row > 1:
            worksheet.auto_filter.ref = worksheet.dimensions

        worksheet.freeze_panes = "A2"

        ResultadoImportacaoXlsxService.__ajustar_larguras(
            worksheet,
            headers)

        ResultadoImportacaoXlsxService.__ajustar_alturas(
            worksheet)

        arquivo = BytesIO()
        workbook.save(arquivo)
        arquivo.seek(0)

        return arquivo

    @staticmethod
    def __obter_dados_linha(linha: dict[str, Any]) -> dict[str, Any]:
        return {
            "CPF": linha.get("CPF", ""),
            "Nome": linha.get("Nome", ""),
            "Empresa": linha.get("Empresa", ""),
            "Placa": linha.get("Placa", ""),
            "De": linha.get("De", ""),
            "Até": linha.get("Até", ""),
            "Resultado": linha.get("Resultado", ""),
        }
    
    @staticmethod
    def __obter_erros(linha: dict[str, Any]) -> list[str]:
        erros = []

        for erro in linha.get("erros", []):
            if isinstance(erro, dict):
                mensagem = erro.get("mensagem")
            else:
                mensagem = getattr(erro, "mensagem", None)

            if mensagem:
                erros.append(str(mensagem))

        return erros

    @staticmethod
    def __obter_headers(linhas: list[dict[str, Any]]) -> list[str]:
        headers = []
        headers_vistos = set()

        for dados in linhas:
            for chave in dados.keys():
                if chave not in headers_vistos:
                    headers.append(chave)
                    headers_vistos.add(chave)

        return headers

    @staticmethod
    def __ajustar_larguras(worksheet, headers: list[str]) -> None:
        for numero_coluna, header in enumerate(headers, start=1):
            letra = get_column_letter(numero_coluna)

            if header == "Erros":
                worksheet.column_dimensions[letra].width = 48
                continue

            maior_tamanho = len(str(header))

            for numero_linha in range(2, worksheet.max_row + 1):
                valor = worksheet.cell(
                    row=numero_linha,
                    column=numero_coluna).value

                if valor is not None:
                    maior_tamanho = max(
                        maior_tamanho,
                        len(str(valor)))

            worksheet.column_dimensions[letra].width = min(
                max(maior_tamanho + 2, 12), 35)

    @staticmethod
    def __ajustar_alturas(worksheet) -> None:
        for numero_linha in range(2, worksheet.max_row + 1):
            maior_quantidade_linhas = 1

            for numero_coluna in range(1, worksheet.max_column + 1):
                valor = worksheet.cell(
                    row=numero_linha,
                    column=numero_coluna).value

                if valor:
                    quantidade = str(valor).count("\n") + 1
                    maior_quantidade_linhas = max(
                        maior_quantidade_linhas,
                        quantidade)

            worksheet.row_dimensions[numero_linha].height = (
                maior_quantidade_linhas * 18)
